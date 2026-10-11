"""Fichier de validation des traductions (xlsx) : français / langue proposée / correction / remarque / statut.
Entrées : i18n/fr.json, i18n/xx.json, et tools/i18n/work/xx/serveur.json (astuces et rappels, voir SKILL).
Usage : python3 tools/i18n/review.py xx "Nom de la langue en français" sortie.xlsx"""
import json, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
LG, NAME, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
fr = json.load(open(f"{ROOT}/i18n/fr.json", encoding="utf-8")); xx = json.load(open(f"{ROOT}/i18n/{LG}.json", encoding="utf-8"))
SRV = json.load(open(f"{ROOT}/tools/i18n/work/{LG}/serveur.json", encoding="utf-8"))   # [{"type","n","fr","xx"}]
SCREEN = {"common": "Général", "days": "Jours", "day": "Jours", "week": "Semaine et fuseau", "moment": "Moments de la journée", "places": "Lieux",
  "emoGroup": "Choix d'icônes", "err": "Messages d'erreur", "error": "Messages d'erreur", "srv": "Messages d'erreur", "garden": "Jardin",
  "oneoff": "Routines exceptionnelles", "tpl": "Modèles", "badge": "Badges", "tier": "Niveaux de récompenses", "effort": "Récompenses",
  "reward": "Récompenses", "help": "Général", "btn": "Boutons", "auth": "Connexion", "pair": "Téléphone de l'enfant", "device": "Téléphone de l'enfant",
  "onboard": "Inscription", "relations": "Membres", "member": "Membres", "role": "Membres", "sep": "Général", "offline": "Général", "nav": "Général",
  "meta": "Général", "dialog": "Général", "tab": "Onglets parent", "parent": "Parent : aujourd'hui", "child": "Enfants", "routine": "Routines",
  "validate": "Validation", "family": "Famille", "pause": "Pauses", "push": "Rappels", "print": "Impression", "pin": "Code parent",
  "backup": "Sauvegarde", "account": "Compte", "malus": "Étoiles cassées", "ctab": "Mode enfant", "kid": "Mode enfant", "bloom": "Mode enfant",
  "timer": "Mode enfant", "form": "Formulaires", "lang": "Langue"}
PL = {"zero", "one", "two", "few", "many", "other"}
def rows(f, x, p=""):
    if isinstance(x, dict) and x and set(x) <= PL and not isinstance(f, dict): f = {"other": f}
    if isinstance(f, dict) and f and set(f) <= PL:   # pluriel : une ligne par forme de la langue cible
        xd = x if isinstance(x, dict) else {"other": x}
        for c, v in xd.items(): yield f"{p} [{c}]", f.get(c, f.get("other")), v
    elif isinstance(f, dict):
        for k in f: yield from rows(f[k], (x or {}).get(k) if isinstance(x, dict) else None, f"{p}.{k}" if p else k)
    elif isinstance(f, list):
        for i, v in enumerate(f): yield f"{p}[{i + 1}]", v, x[i] if isinstance(x, list) and i < len(x) else None
    else: yield p, f, x
wb = Workbook(); ws = wb.active; ws.title = "Relecture app"
yellow = PatternFill("solid", fgColor="FFF2CC"); head = PatternFill("solid", fgColor="DDF2E5"); wrap = Alignment(wrap_text=True, vertical="top")
ws["A1"] = f"SoisFier : relecture de la traduction {NAME}"; ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = f"Lire la colonne « {NAME} proposé ». Si une phrase ne convient pas, écrire la bonne version dans « Correction » (colonne jaune) ; laisser vide si c'est correct. Indiquer le statut."
ws["A3"] = "{nom} = valeur remplacée par l'application (ne pas traduire). [one], [few]… = forme utilisée selon le nombre."
COLS = ["Écran", "Clé (technique)", "Français", f"{NAME} proposé", "Correction", "Remarque", "Statut"]
def sheet_head(s, r):
    for i, c in enumerate(COLS, 1):
        cell = s.cell(r, i, c); cell.font = Font(bold=True); cell.fill = head
    for col, w in zip("ABCDEFG", (20, 26, 48, 48, 40, 30, 12)): s.column_dimensions[col].width = w
dv = DataValidation(type="list", formula1='"à relire,validé,corrigé"', allow_blank=True); ws.add_data_validation(dv)
sheet_head(ws, 5); r = 6
for k, f, x in rows(fr, xx):
    ws.append([SCREEN.get(k.split(".")[0].split("[")[0], "Général"), k, f, x, None, None, "à relire"])
    for c in range(1, 8): ws.cell(r, c).alignment = wrap
    ws.cell(r, 5).fill = yellow; dv.add(ws.cell(r, 7)); r += 1
ws.freeze_panes = "A6"; ws.auto_filter.ref = f"A5:G{r - 1}"
s2 = wb.create_sheet("Astuces et rappels"); s2.append(["Type", "N°", "Français", f"{NAME} proposé", "Correction", "Remarque", "Statut"])
for c in s2[1]: c.font = Font(bold=True); c.fill = head
for col, w in zip("ABCDEFG", (22, 6, 55, 55, 45, 30, 12)): s2.column_dimensions[col].width = w
dv2 = DataValidation(type="list", formula1='"à relire,validé,corrigé"', allow_blank=True); s2.add_data_validation(dv2)
for i, it in enumerate(SRV, 2):
    s2.append([it["type"], it["n"], it["fr"], it["xx"], None, None, "à relire"])
    for c in range(1, 8): s2.cell(i, c).alignment = wrap
    s2.cell(i, 5).fill = yellow; dv2.add(s2.cell(i, 7))
s2.freeze_panes = "A2"
s3 = wb.create_sheet("Mode d'emploi")
s3.append([f"Relire le mode d'emploi en ligne : guide/{LG}/ (lien de langue en haut du guide français). Noter ici les corrections : partie, phrase actuelle, correction."])
s3.append(["Partie", "Phrase actuelle", "Correction", "Statut"])
for c in s3[2]: c.font = Font(bold=True); c.fill = head
for col, w in zip("ABCD", (16, 60, 60, 12)): s3.column_dimensions[col].width = w
wb.save(OUT); print(OUT, r - 6, "textes app,", len(SRV), "astuces et rappels")
